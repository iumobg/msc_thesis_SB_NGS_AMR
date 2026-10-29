"""
'wastewater metagenome' örneklerinde hangi metadata alanları ne sıklıkla dolu?
Üç kaynak: NCBI BioSample, NCBI SRA, EBI ENA

Çıktılar:
  - Ekranda her kaynak için özet + en sık dolu alanlar (yüzde çubuklu)
  - Kaynak karşılaştırma tablosu ve örnek örtüşmesi (aynı örnek kaç kaynakta)
  - metadata_raporu.xlsx  (pandas + openpyxl varsa; yoksa CSV dosyaları)
  - metadata_karsilastirma.png (matplotlib varsa)
"""
import time
import json
import statistics
import urllib.request
import urllib.parse
import urllib.error
import xml.etree.ElementTree as ET
from collections import Counter, defaultdict

from Bio import Entrez

import ssl
ssl._create_default_https_context = ssl._create_unverified_context

# ------------------------------------------------------------------ AYARLAR
Entrez.email = "boysansebnem@gmail.com"        # NCBI için zorunlu
# Entrez.api_key = "..."                       # varsa: hız limiti 3 -> 10 istek/sn

ORGANISM = "wastewater metagenome"
TAXID = 527639                                  # NCBI Taxonomy: wastewater metagenome
NCBI_QUERY = f'"{ORGANISM}"[Organism]'          # istersen: OR "sludge metagenome"[Organism]

MAX_RECORDS = 2000      # her kaynak için üst sınır (önce küçük başla)
BATCH = 200             # NCBI efetch parti büyüklüğü
ENA_BATCH = 100         # ENA XML parti büyüklüğü
SLEEP = 0.4             # istekler arası bekleme (sn)
TOP_N = 25              # ekranda gösterilecek alan sayısı

RUN_BIOSAMPLE = True
RUN_SRA = True
RUN_ENA = True

EMPTY = {"", "missing", "not applicable", "not collected", "not provided",
         "unknown", "na", "n/a", "none", "restricted access", "not determined",
         "-", ".", "null", "nan", "missing: not provided", "missing: not collected",
         "missing: not applicable", "missing: restricted access"}

# Sistemin otomatik eklediği, kullanıcının girmediği alanlar -> sayılmaz
SYSTEM_PREFIXES = ("ena_", "insdc_")
SYSTEM_FIELDS = {"submitter_id", "external_id", "sra_accession", "biosamplemodel",
                 "ena_checklist"}

# Aynı anlamdaki alan adlarını birleştir (ENA/MIxS etiketleri -> NCBI harmonized adları)
ALIASES = {
    "geographic_location_(country_and/or_sea)": "geo_loc_name",
    "geographic_location_(region_and_locality)": "geo_loc_name_region",
    "geographic_location_(latitude)": "lat_lon",
    "geographic_location_(longitude)": "lat_lon",
    "latitude_and_longitude": "lat_lon",
    "latitude": "lat_lon",
    "longitude": "lat_lon",
    "geographic_location_(depth)": "depth",
    "geographic_location_(elevation)": "elev",
    "elevation": "elev",
    "collection_date": "collection_date",
    "sample_collection_date": "collection_date",
    "broad_scale_environmental_context": "env_broad_scale",
    "broad-scale_environmental_context": "env_broad_scale",
    "environment_(biome)": "env_broad_scale",
    "local_environmental_context": "env_local_scale",
    "environment_(feature)": "env_local_scale",
    "environmental_medium": "env_medium",
    "environment_(material)": "env_medium",
    "sequencing_method": "seq_meth",
    "investigation_type": "investigation_type",
    "project_name": "project_name",
    "isolation_source": "isolation_source",
    "sample_name": "sample_name",
    "host_scientific_name": "host",
    "temperature": "temp",
    "wastewater_type": "wastewater_type",
    "sewage_type": "sewage_type",
}

# Sık görülen alanların Türkçe açıklamaları (sözlükte olmayanlar NCBI'ın resmi
# İngilizce tanımından gelir; o da yoksa boş kalır)
FIELD_DESC_TR = {
    # --- temel / kimlik
    "sample_name": "Örneğe gönderen tarafından verilen ad",
    "isolate": "Örneğe/izolata verilen tanımlayıcı ad",
    "strain": "Suş adı (metagenomlarda genellikle örnek adı yerine kullanılır)",
    "description": "Örnek hakkında serbest metin açıklama",
    "project_name": "Örneğin ait olduğu projenin adı",
    "investigation_type": "Çalışma türü (ör. metagenome, mimarks-survey)",
    "env_package": "Kullanılan MIxS çevre paketi (ör. wastewater/sludge)",
    "sample_type": "Örneğin tipi (ör. metagenomic assembly, environmental sample)",
    "source_material_id": "Kaynak materyalin kimlik numarası (biyobanka, laboratuvar no.)",
    "host_subject_id": "Konak bireyin anonim kimlik numarası",
    # --- zaman ve yer
    "collection_date": "Örneğin toplandığı tarih",
    "geo_loc_name": "Örneğin alındığı ülke ve bölge (ör. 'USA: California')",
    "geo_loc_name_region": "Bölge ve yerleşim yeri (ENA'da ayrı alan)",
    "lat_lon": "Örnekleme noktasının enlem-boylam koordinatları",
    "depth": "Örnekleme derinliği",
    "elev": "Rakım (deniz seviyesinden yükseklik)",
    # --- çevre (ENVO terimleri)
    "env_broad_scale": "Geniş ölçekli çevre/biyom (ör. 'urban biome')",
    "env_local_scale": "Örneğin alındığı yerel ortam (ör. 'wastewater treatment plant')",
    "env_medium": "Örneğin fiziksel materyali (ör. 'wastewater', 'activated sludge')",
    "isolation_source": "Örneğin alındığı yer/kaynak, serbest metin (ör. 'influent')",
    "host": "Konak organizma (ör. Homo sapiens)",
    "rel_to_oxygen": "Oksijenle ilişki (aerobik / anaerobik)",
    # --- örnekleme ve işleme
    "collected_by": "Örneği toplayan kişi veya kurum",
    "biomaterial_provider": "Örneği sağlayan kişi veya kurum",
    "samp_collect_device": "Örnek toplama aleti/yöntemi (ör. otomatik kompozit örnekleyici)",
    "samp_size": "Toplanan örnek miktarı (hacim veya kütle)",
    "samp_mat_process": "Örneğe uygulanan işlem (filtreleme, konsantrasyon vb.)",
    "samp_store_temp": "Örneğin saklandığı sıcaklık",
    "samp_store_dur": "Örneğin saklandığı süre",
    "samp_vol_we_dna_ext": "DNA ekstraksiyonu için kullanılan örnek hacmi/ağırlığı",
    "nucl_acid_ext": "Nükleik asit (DNA/RNA) ekstraksiyon yöntemi veya kiti",
    "nucl_acid_amp": "Nükleik asit çoğaltma yöntemi",
    "seq_meth": "Dizileme yöntemi / cihazı",
    "purpose_of_sampling": "Örnekleme amacı (ör. halk sağlığı sürveyansı)",
    "purpose_of_sequencing": "Dizileme amacı",
    # --- atık su sürveyansı (NCBI wastewater paketi)
    "ww_population": "Örneklenen atık su ağının hizmet verdiği nüfus",
    "ww_sample_type": "Örnek tipi: kompozit (zamana yayılmış) veya anlık (grab)",
    "ww_sample_duration": "Kompozit örneğin toplandığı süre (ör. 24 saat)",
    "ww_sample_site": "Örnekleme noktası (arıtma tesisi girişi, kanalizasyon hattı vb.)",
    "ww_sample_matrix": "Örnek matrisi (ham atık su, birincil çamur vb.)",
    "ww_flow": "Örnekleme sırasındaki atık su debisi",
    "ww_processing_protocol": "Atık su örneğini işleme protokolü",
    "ww_endog_control_1": "Normalizasyon için endojen kontrol (ör. PMMoV)",
    "ww_surv_target_1": "Sürveyansta aranan hedef organizma (ör. SARS-CoV-2)",
    "ww_surv_target_1_known_present": "Hedef organizmanın örnekte bulunup bulunmadığı",
    "wastewater_type": "Atık su tipi (evsel, endüstriyel, hastane vb.)",
    "sewage_type": "Kanalizasyon/atık su tipi",
    "industrial_effluent_percent": "Atık suyun yüzde kaçının endüstriyel kaynaklı olduğu",
    "sludge_retent_time": "Çamur bekletme süresi (arıtma tesisinde)",
    # --- fiziko-kimyasal ölçümler
    "temp": "Örnek sıcaklığı",
    "ph": "pH değeri",
    "salinity": "Tuzluluk",
    "alkalinity": "Alkalinite",
    "diss_oxygen": "Çözünmüş oksijen",
    "turbidity": "Bulanıklık",
    "cod": "Kimyasal oksijen ihtiyacı (KOİ)",
    "bod": "Biyokimyasal oksijen ihtiyacı (BOİ)",
    "tot_org_carb": "Toplam organik karbon",
    "tot_nitro": "Toplam azot",
    "tot_phosp": "Toplam fosfor",
    "ammonium": "Amonyum konsantrasyonu",
    "nitrate": "Nitrat konsantrasyonu",
    "phosphate": "Fosfat konsantrasyonu",
    "suspend_solids": "Askıda katı madde",
    # --- SRA teknik alanları
    "lib:library_strategy": "Dizileme stratejisi (WGS, AMPLICON, RNA-Seq vb.)",
    "lib:library_source": "Kütüphane kaynağı (METAGENOMIC, METATRANSCRIPTOMIC vb.)",
    "lib:library_selection": "Kütüphane seçim yöntemi (RANDOM, PCR vb.)",
    "lib:library_construction_protocol": "Kütüphane hazırlama protokolünün açıklaması",
    "lib:instrument_model": "Dizileme cihazı modeli (ör. Illumina NovaSeq 6000)",
}

NCBI_ATTR_URL = "https://www.ncbi.nlm.nih.gov/biosample/docs/attributes/?format=xml"
NCBI_DESC = {}        # harmonized_name -> resmi İngilizce açıklama
NCBI_SYNONYMS = {}    # eşanlamlı ad -> harmonized_name


# ------------------------------------------------------------------ YARDIMCILAR
def norm_field(name):
    """Alan adını standartlaştır: küçük harf, boşluk/tire -> alt çizgi, eşanlamlıları birleştir."""
    n = _basic_norm(name)
    n = ALIASES.get(n, n)
    return NCBI_SYNONYMS.get(n, n)


def _basic_norm(name):
    n = (name or "").strip().lower().replace(" ", "_").replace("-", "_")
    while "__" in n:
        n = n.replace("__", "_")
    return n


def load_ncbi_dictionary():
    """NCBI'ın resmi BioSample alan listesini indir: açıklamalar + eşanlamlılar."""
    try:
        root = ET.fromstring(http_get(NCBI_ATTR_URL))
    except Exception as e:
        print(f"(NCBI alan sözlüğü indirilemedi: {e}; yalnızca Türkçe sözlük kullanılacak)")
        return
    for a in root.iter("Attribute"):
        hn = _basic_norm(a.findtext("HarmonizedName"))
        if not hn:
            continue
        desc = " ".join((a.findtext("Description") or "").split())
        if desc:
            NCBI_DESC[hn] = desc
        for syn in [a.findtext("Name")] + [s.text for s in a.iter("Synonym")]:
            s = _basic_norm(syn)
            if s and s != hn:
                NCBI_SYNONYMS.setdefault(s, hn)
    print(f"NCBI alan sözlüğü yüklendi: {len(NCBI_DESC)} alan tanımı")


def describe(field, max_len=None):
    d = FIELD_DESC_TR.get(field) or NCBI_DESC.get(field) or ""
    if max_len and len(d) > max_len:
        d = d[:max_len - 1].rstrip() + "…"
    return d


def is_filled(value):
    return (value or "").strip().lower() not in EMPTY


def is_system(field):
    return field in SYSTEM_FIELDS or field.startswith(SYSTEM_PREFIXES)


def http_get(url, data=None, tries=4):
    """Basit GET/POST, hata olursa artan beklemeyle tekrar dener."""
    for k in range(tries):
        try:
            body = urllib.parse.urlencode(data).encode() if data else None
            with urllib.request.urlopen(url, data=body, timeout=120) as r:
                return r.read()
        except (urllib.error.URLError, TimeoutError) as e:
            if k == tries - 1:
                raise
            print(f"   ! istek hatası ({e}), tekrar deneniyor...")
            time.sleep(2 * (k + 1))


def make_record(sample_id, biosample, project, fields, checklist=None):
    """Her kaynak için ortak örnek kaydı formatı."""
    return {"id": sample_id, "biosample": biosample, "project": project,
            "fields": fields, "checklist": checklist}


# ------------------------------------------------------------------ 1) NCBI BioSample
def fetch_biosample():
    print("\n[BioSample] aranıyor...")
    h = Entrez.esearch(db="biosample", term=NCBI_QUERY, retmax=MAX_RECORDS)
    r = Entrez.read(h)
    ids = r["IdList"]
    print(f"   toplam {r['Count']} kayıt var, {len(ids)} tanesi indirilecek")

    records = []
    for i in range(0, len(ids), BATCH):
        h = Entrez.efetch(db="biosample", id=",".join(ids[i:i + BATCH]),
                          rettype="full", retmode="xml")
        records += parse_biosample_xml(h.read())
        print(f"   {min(i + BATCH, len(ids))}/{len(ids)}")
        time.sleep(SLEEP)
    return records


def parse_biosample_xml(xml_bytes):
    out = []
    root = ET.fromstring(xml_bytes)
    for bs in root.iter("BioSample"):
        acc = bs.get("accession")
        link = bs.find(".//Link[@target='bioproject']")
        proj = link.get("label") if link is not None else None
        pkg = bs.findtext("Package")
        fields = set()
        for a in bs.iter("Attribute"):
            name = norm_field(a.get("harmonized_name") or a.get("attribute_name"))
            if name and not is_system(name) and is_filled(a.text):
                fields.add(name)
        out.append(make_record(acc, acc, proj, fields, pkg))
    return out


# ------------------------------------------------------------------ 2) NCBI SRA
SRA_TECH = {  # SRA'da deney (experiment) düzeyindeki teknik alanlar
    "lib:library_strategy": ".//LIBRARY_DESCRIPTOR/LIBRARY_STRATEGY",
    "lib:library_source": ".//LIBRARY_DESCRIPTOR/LIBRARY_SOURCE",
    "lib:library_selection": ".//LIBRARY_DESCRIPTOR/LIBRARY_SELECTION",
    "lib:library_construction_protocol": ".//LIBRARY_DESCRIPTOR/LIBRARY_CONSTRUCTION_PROTOCOL",
    "lib:instrument_model": ".//PLATFORM//INSTRUMENT_MODEL",
}


def fetch_sra():
    print("\n[SRA] aranıyor...")
    h = Entrez.esearch(db="sra", term=NCBI_QUERY, retmax=0, usehistory="y")
    r = Entrez.read(h)
    total = int(r["Count"])
    n = min(total, MAX_RECORDS)
    print(f"   toplam {total} SRA deneyi var, {n} tanesi indirilecek")

    by_sample = {}  # aynı örneğin birden fazla deneyi olabilir -> örnek bazında birleştir
    tech = Counter()
    n_exp = 0
    for start in range(0, n, BATCH):
        h = Entrez.efetch(db="sra", WebEnv=r["WebEnv"], query_key=r["QueryKey"],
                          retstart=start, retmax=min(BATCH, n - start), retmode="xml")
        exps, t = parse_sra_xml(h.read())
        n_exp += len(exps)
        tech.update(t)
        for rec in exps:
            if rec["id"] in by_sample:
                by_sample[rec["id"]]["fields"] |= rec["fields"]
            else:
                by_sample[rec["id"]] = rec
        print(f"   {min(start + BATCH, n)}/{n}")
        time.sleep(SLEEP)
    print(f"   {n_exp} deney -> {len(by_sample)} benzersiz örnek")
    return list(by_sample.values()), tech, n_exp


def parse_sra_xml(xml_bytes):
    out, tech = [], Counter()
    root = ET.fromstring(xml_bytes)
    for pkg in root.iter("EXPERIMENT_PACKAGE"):
        sample = pkg.find("SAMPLE")
        if sample is None:
            continue
        sra_acc = sample.get("accession")
        biosample = sample.findtext(".//EXTERNAL_ID[@namespace='BioSample']")
        proj = pkg.findtext(".//STUDY//EXTERNAL_ID[@namespace='BioProject']")
        fields = set()
        for a in sample.iter("SAMPLE_ATTRIBUTE"):
            name = norm_field(a.findtext("TAG"))
            if name and not is_system(name) and is_filled(a.findtext("VALUE")):
                fields.add(name)
        for key, path in SRA_TECH.items():
            if is_filled(pkg.findtext(path)):
                tech[key] += 1
        out.append(make_record(sra_acc, biosample, proj, fields))
    return out, tech


# ------------------------------------------------------------------ 3) EBI ENA
ENA_PORTAL = "https://www.ebi.ac.uk/ena/portal/api/search"
ENA_XML = "https://www.ebi.ac.uk/ena/browser/api/xml/"


def fetch_ena():
    print("\n[ENA] aranıyor...")
    raw = http_get(ENA_PORTAL, data={
        "result": "sample",
        "query": f"tax_eq({TAXID})",
        "fields": "sample_accession,secondary_sample_accession",
        "format": "json",
        "limit": MAX_RECORDS,
    })
    rows = json.loads(raw) if raw.strip() else []
    accs = [r["sample_accession"] for r in rows if r.get("sample_accession")]
    print(f"   {len(accs)} örnek indirilecek")

    records = []
    for i in range(0, len(accs), ENA_BATCH):
        xml_bytes = http_get(ENA_XML + ",".join(accs[i:i + ENA_BATCH]))
        records += parse_ena_xml(xml_bytes)
        print(f"   {min(i + ENA_BATCH, len(accs))}/{len(accs)}")
        time.sleep(SLEEP)
    return records


def parse_ena_xml(xml_bytes):
    out = []
    root = ET.fromstring(xml_bytes)
    for s in root.iter("SAMPLE"):
        biosample = s.findtext(".//EXTERNAL_ID[@namespace='BioSample']") or s.get("accession")
        proj = None
        for x in s.iter("XREF_LINK"):
            if (x.findtext("DB") or "").upper() in ("ENA-STUDY", "BIOPROJECT"):
                proj = x.findtext("ID")
                break
        fields, checklist = set(), None
        for a in s.iter("SAMPLE_ATTRIBUTE"):
            raw_tag = a.findtext("TAG") or ""
            if raw_tag.upper() == "ENA-CHECKLIST":
                checklist = a.findtext("VALUE")
            name = norm_field(raw_tag)
            if name and not is_system(name) and is_filled(a.findtext("VALUE")):
                fields.add(name)
        out.append(make_record(s.get("accession"), biosample, proj, fields, checklist))
    return out


# ------------------------------------------------------------------ ANALİZ
def summarize(records):
    """Alan -> (dolu örnek sayısı, yüzde, proje sayısı) tablosu ve genel özet."""
    n = len(records)
    present, projects = Counter(), defaultdict(set)
    for r in records:
        for f in r["fields"]:
            present[f] += 1
            if r["project"]:
                projects[f].add(r["project"])
    rows = [{"alan": f, "dolu_ornek": c, "yuzde": round(100 * c / n, 1) if n else 0,
             "proje_sayisi": len(projects[f]), "aciklama": describe(f)}
            for f, c in present.most_common()]
    per_sample = [len(r["fields"]) for r in records] or [0]
    summary = {
        "ornek": n,
        "proje": len({r["project"] for r in records if r["project"]}),
        "farkli_alan": len(present),
        "ornek_basi_ort_alan": round(statistics.mean(per_sample), 1),
        "ornek_basi_medyan_alan": statistics.median(per_sample),
        "alan_90+": sum(1 for r in rows if r["yuzde"] >= 90),
        "alan_50-90": sum(1 for r in rows if 50 <= r["yuzde"] < 90),
        "alan_10-50": sum(1 for r in rows if 10 <= r["yuzde"] < 50),
        "alan_<10": sum(1 for r in rows if r["yuzde"] < 10),
    }
    checklists = Counter(r["checklist"] for r in records if r.get("checklist"))
    return rows, summary, checklists


def bar(pct, width=25):
    full = int(round(pct / 100 * width))
    return "█" * full + "·" * (width - full)


def print_source(name, rows, summary, checklists, extra=None):
    line = "=" * 78
    print(f"\n{line}\n {name}\n{line}")
    print(f" Örnek sayısı            : {summary['ornek']}")
    print(f" BioProject / çalışma    : {summary['proje']}")
    print(f" Farklı metadata alanı   : {summary['farkli_alan']}")
    print(f" Örnek başına dolu alan  : ortalama {summary['ornek_basi_ort_alan']}, "
          f"medyan {summary['ornek_basi_medyan_alan']}")
    print(f" Doluluk dağılımı        : ≥%90: {summary['alan_90+']} alan | "
          f"%50-90: {summary['alan_50-90']} | %10-50: {summary['alan_10-50']} | "
          f"<%10: {summary['alan_<10']}")
    if checklists:
        top = ", ".join(f"{k} ({v})" for k, v in checklists.most_common(3))
        print(f" En sık paket/checklist  : {top}")
    if extra:
        print(extra)
    print(f"\n En sık dolu {min(TOP_N, len(rows))} alan:")
    print(f" {'alan':<32} {'doluluk':<27} {'%':>6} {'örnek':>7} {'proje':>6}")
    print(" " + "-" * 76)
    for r in rows[:TOP_N]:
        print(f" {r['alan'][:32]:<32} {bar(r['yuzde'])}  {r['yuzde']:>5.1f} "
              f"{r['dolu_ornek']:>7} {r['proje_sayisi']:>6}")
        d = describe(r["alan"], max_len=90)
        if d:
            print(f"   ↳ {d}")
    print("\n Not: 'proje' sütunu düşükse (ör. 1-2), yüksek yüzde tek bir büyük "
          "projeden geliyor olabilir.")


def compare(results):
    """Kaynaklar arası yan yana tablo: her alan için her kaynaktaki yüzde."""
    sources = list(results)
    pct = {s: {r["alan"]: r["yuzde"] for r in results[s][0]} for s in sources}
    all_fields = set().union(*[set(p) for p in pct.values()])
    table = []
    for f in all_fields:
        vals = {s: pct[s].get(f, 0.0) for s in sources}
        table.append({"alan": f, **{f"{s}_%": vals[s] for s in sources},
                      "ortalama_%": round(sum(vals.values()) / len(sources), 1),
                      "kac_kaynakta": sum(1 for v in vals.values() if v > 0),
                      "aciklama": describe(f)})
    table.sort(key=lambda r: -r["ortalama_%"])
    return table


def print_compare(table, sources):
    print("\n" + "=" * 78 + "\n KAYNAK KARŞILAŞTIRMASI (örneklerin yüzde kaçında dolu)\n" + "=" * 78)
    head = f" {'alan':<32}" + "".join(f"{s:>12}" for s in sources)
    print(head + "\n " + "-" * (len(head) - 1))
    for r in table[:TOP_N]:
        print(f" {r['alan'][:32]:<32}" + "".join(f"{r[s + '_%']:>11.1f}%" for s in sources)
              + f"   {describe(r['alan'], max_len=45)}")
    common = [r["alan"] for r in table if r["kac_kaynakta"] == len(sources)]
    print(f"\n Tüm kaynaklarda ortak alan sayısı: {len(common)}")
    core = [r["alan"] for r in table if all(r[s + "_%"] >= 80 for s in sources)]
    if core:
        print(f" Her kaynakta ≥%80 dolu 'çekirdek' alanlar: {', '.join(core)}")


def print_overlap(results):
    sets = {s: {r["biosample"] for r in results[s][3] if r["biosample"]} for s in results}
    if len(sets) < 2:
        return
    print("\n" + "=" * 78 + "\n ÖRNEK ÖRTÜŞMESİ (BioSample numarasına göre)\n" + "=" * 78)
    names = list(sets)
    for i, a in enumerate(names):
        for b in names[i + 1:]:
            print(f" {a} ∩ {b}: {len(sets[a] & sets[b])} ortak örnek")
    if len(sets) == 3:
        print(f" Üçünde birden: {len(set.intersection(*sets.values()))}")
    for s in names:
        others = set().union(*[sets[o] for o in names if o != s])
        print(f" Sadece {s}'da: {len(sets[s] - others)}")
    print(" Not: NCBI ve ENA (INSDC) verileri birbirine aynalanır; örneklerin çoğu "
          "ortak olabilir.\n Farklar genellikle MAX_RECORDS sınırından ya da "
          "sorgu kapsamından gelir.")


# ------------------------------------------------------------------ EXCEL / GRAFİK
# Renk paleti: tek renk ailesi (mavi), renk körlüğüne uygun, yazdırınca da okunur
C_HEADER = "1F3A5F"       # başlık zemini (koyu lacivert)
C_SUBHEAD = "D9E2EC"      # alt başlık / bölüm zemini
C_BORDER = "BFC9D4"
SCALE = [(0, "FFFFFF"), (50, "9ECAE1"), (100, "3182BD")]   # % renk ölçeği
CLASS_BANDS = [            # doluluk sınıfları (ölçekle aynı renk ailesi)
    (90, "Yüksek (≥%90)", "3182BD", "FFFFFF"),
    (50, "Orta (%50–90)", "9ECAE1", "000000"),
    (10, "Düşük (%10–50)", "DEEBF7", "000000"),
    (0, "Çok düşük (<%10)", "FFFFFF", "7F7F7F"),
]

SUMMARY_LABELS = {
    "ornek": "Örnek sayısı",
    "proje": "BioProject / çalışma sayısı",
    "farkli_alan": "Farklı metadata alanı sayısı",
    "ornek_basi_ort_alan": "Örnek başına dolu alan (ortalama)",
    "ornek_basi_medyan_alan": "Örnek başına dolu alan (medyan)",
    "alan_90+": "Doluluğu ≥%90 olan alan sayısı",
    "alan_50-90": "Doluluğu %50–90 olan alan sayısı",
    "alan_10-50": "Doluluğu %10–50 olan alan sayısı",
    "alan_<10": "Doluluğu <%10 olan alan sayısı",
}

COLUMN_GUIDE = [
    ("alan", "Metadata alanının standartlaştırılmış adı (NCBI harmonized name veya eşdeğeri)."),
    ("aciklama", "Alanın anlamı. Türkçe sözlükten ya da NCBI resmi tanımından gelir."),
    ("dolu_ornek", "Bu alanın geçerli bir değer içerdiği örnek sayısı."),
    ("yuzde", "dolu_ornek / toplam örnek × 100. Alanın o kaynaktaki doluluk oranı."),
    ("proje_sayisi", "Alanı dolduran farklı BioProject sayısı. Düşükse yüksek yüzde tek bir "
                     "büyük projeden kaynaklanıyor olabilir (proje yanlılığı)."),
    ("doluluk_sinifi", "Yüzdeye göre sınıf: Yüksek ≥90, Orta 50–90, Düşük 10–50, Çok düşük <10."),
    ("<Kaynak>_%", "Karşılaştırma sayfası: alanın o kaynaktaki örneklerin yüzde kaçında dolu olduğu. "
                   "%0, alanın o kaynakta bu adla hiç bulunmadığı anlamına gelir."),
    ("ortalama_%", "Kaynak yüzdelerinin basit (ağırlıksız) ortalaması. Sıralama ölçütüdür."),
    ("kac_kaynakta", "Alanın en az bir örnekte dolu olduğu kaynak sayısı (1–3)."),
    ("dolu_deney", "SRA_teknik: alanın dolu olduğu SRA deneyi (experiment) sayısı."),
]


def doluluk_sinifi(pct):
    for lo, label, _, _ in CLASS_BANDS:
        if pct >= lo:
            return label
    return CLASS_BANDS[-1][1]


def make_chart(table, sources):
    try:
        import matplotlib
        matplotlib.use("Agg")
        import matplotlib.pyplot as plt
        top = table[:20][::-1]
        colors = ["#08519C", "#6BAED6", "#BDBDBD"]   # renk körlüğüne uygun mavi/gri tonları
        fig, ax = plt.subplots(figsize=(10, 8))
        h = 0.8 / len(sources)
        for k, s in enumerate(sources):
            ax.barh([i + k * h for i in range(len(top))], [r[s + "_%"] for r in top],
                    height=h, label=s, color=colors[k % len(colors)], edgecolor="white")
        ax.set_yticks([i + h * (len(sources) - 1) / 2 for i in range(len(top))])
        ax.set_yticklabels([r["alan"] for r in top])
        ax.set_xlabel("Alanın dolu olduğu örneklerin oranı (%)")
        ax.set_xlim(0, 100)
        ax.grid(axis="x", color="#E0E0E0", linewidth=0.8)
        ax.set_axisbelow(True)
        for side in ("top", "right"):
            ax.spines[side].set_visible(False)
        ax.set_title(f"'{ORGANISM}' — en sık dolu 20 metadata alanı (kaynaklara göre)")
        ax.legend(loc="lower right", frameon=False)
        fig.tight_layout()
        fig.savefig("metadata_karsilastirma.png", dpi=200)
        plt.close(fig)
        print("✔ metadata_karsilastirma.png kaydedildi")
        return "metadata_karsilastirma.png"
    except Exception as e:
        print(f"(Grafik çizilemedi: {e})")
        return None


def save_outputs(results, table, extras):
    sources = list(results)
    for s in sources:
        for r in results[s][0]:
            r["doluluk_sinifi"] = doluluk_sinifi(r["yuzde"])
    for r in table:
        r["doluluk_sinifi"] = doluluk_sinifi(r["ortalama_%"])
    chart = make_chart(table, sources) if table else None

    try:
        write_excel(results, table, extras, chart)
        print("✔ metadata_raporu.xlsx kaydedildi")
    except Exception as e:
        import csv
        print(f"\n(Excel yazılamadı: {e}) -> CSV olarak kaydediliyor")
        for s in sources:
            with open(f"metadata_{s}.csv", "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=list(results[s][0][0].keys()))
                w.writeheader()
                w.writerows(results[s][0])
        if table:
            with open("metadata_karsilastirma.csv", "w", newline="", encoding="utf-8") as fh:
                w = csv.DictWriter(fh, fieldnames=list(table[0].keys()))
                w.writeheader()
                w.writerows(table)
        print("✔ CSV dosyaları kaydedildi")


def write_excel(results, table, extras, chart):
    from datetime import datetime
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, CellIsRule
    from openpyxl.utils import get_column_letter

    sources = list(results)
    thin = Side(style="thin", color=C_BORDER)
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    f_head = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    f_body = Font(name="Calibri", size=10)
    f_title = Font(name="Calibri", size=14, bold=True, color=C_HEADER)
    f_sec = Font(name="Calibri", size=11, bold=True, color=C_HEADER)
    fill_head = PatternFill("solid", fgColor=C_HEADER)
    fill_sub = PatternFill("solid", fgColor=C_SUBHEAD)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    widths = {"alan": 34, "aciklama": 70, "doluluk_sinifi": 18, "kac_kaynakta": 13,
              "proje_sayisi": 13, "dolu_ornek": 12, "dolu_deney": 12}

    def table_sheet(ws, rows, columns, pct_cols):
        """Başlıklı, kenarlıklı, filtreli tablo + koşullu biçimlendirme."""
        ws.append(columns)
        for r in rows:
            ws.append([r.get(c) for c in columns])
        n = len(rows) + 1
        for j, col in enumerate(columns, 1):
            L = get_column_letter(j)
            ws.column_dimensions[L].width = widths.get(col, 12 if not col.endswith("%") else 11)
            h = ws.cell(1, j)
            h.font, h.fill, h.border = f_head, fill_head, border
            h.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for i in range(2, n + 1):
                c = ws.cell(i, j)
                c.font, c.border = f_body, border
                if col in ("aciklama",):
                    c.alignment = wrap_top
                elif col in pct_cols:
                    c.number_format = "0.0"
                    c.alignment = Alignment(horizontal="center", vertical="top")
                elif isinstance(c.value, (int, float)):
                    c.number_format = "#,##0"
                    c.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    c.alignment = Alignment(vertical="top")
            rng = f"{L}2:{L}{n}"
            if n < 2:
                continue
            if col in pct_cols:          # yüzdeler: beyaz -> mavi ölçek
                ws.conditional_formatting.add(rng, ColorScaleRule(
                    start_type="num", start_value=SCALE[0][0], start_color=SCALE[0][1],
                    mid_type="num", mid_value=SCALE[1][0], mid_color=SCALE[1][1],
                    end_type="num", end_value=SCALE[2][0], end_color=SCALE[2][1]))
            elif col in ("proje_sayisi", "dolu_ornek", "dolu_deney"):   # sayılar: gri çubuk
                ws.conditional_formatting.add(rng, DataBarRule(
                    start_type="num", start_value=0, end_type="max", color="A6A6A6"))
            elif col == "doluluk_sinifi":
                for i in range(2, n + 1):
                    c = ws.cell(i, j)
                    for _, label, bg, fg in CLASS_BANDS:
                        if c.value == label:
                            c.fill = PatternFill("solid", fgColor=bg)
                            c.font = Font(name="Calibri", size=10, color=fg,
                                          bold=label.startswith("Yüksek"))
            elif col == "kac_kaynakta":
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=[str(len(sources))],
                    font=Font(bold=True, color=C_HEADER), fill=fill_sub))
                ws.conditional_formatting.add(rng, CellIsRule(
                    operator="equal", formula=["1"], font=Font(italic=True, color="7F7F7F")))
        ws.freeze_panes = "B2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{n}"
        ws.row_dimensions[1].height = 32
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"

    wb = Workbook()

    # ---- 1) Okuma kılavuzu (yöntem + sütunlar + renk anahtarı)
    ws = wb.active
    ws.title = "Okuma_Kilavuzu"
    ws.column_dimensions["A"].width = 30
    ws.column_dimensions["B"].width = 95
    ws["A1"] = f"Metadata doluluk analizi — '{ORGANISM}'"
    ws["A1"].font = f_title
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    row = 3

    def section(title):
        nonlocal row
        ws.cell(row, 1, title).font = f_sec
        ws.cell(row, 1).fill = fill_sub
        ws.cell(row, 2).fill = fill_sub
        row += 1

    def kv(k, v):
        nonlocal row
        a, b = ws.cell(row, 1, k), ws.cell(row, 2, v)
        a.font = Font(name="Calibri", size=10, bold=True)
        b.font = f_body
        a.alignment, b.alignment = wrap_top, wrap_top
        a.border = b.border = border
        row += 1

    section("1. Yöntem ve tekrarlanabilirlik")
    kv("Veri çekme tarihi", datetime.now().strftime("%Y-%m-%d %H:%M"))
    kv("Organizma / Taxon ID", f"{ORGANISM} (NCBI Taxonomy {TAXID})")
    kv("NCBI sorgusu", NCBI_QUERY + "  (BioSample ve SRA veritabanları, Entrez E-utilities)")
    kv("ENA sorgusu", f"Portal API: result=sample, query=tax_eq({TAXID}); "
                      "öznitelikler Browser API XML'inden")
    kv("Kaynak başına üst sınır", f"MAX_RECORDS = {MAX_RECORDS}")
    kv("Analiz birimi", "Örnek (sample). SRA'da aynı örneğe ait birden fazla deney tek örnekte "
                        "birleştirildi.")
    kv("'Dolu' tanımı", "Değer boş değilse ve şu ifadelerden biri değilse: "
                        + ", ".join(sorted(e for e in EMPTY if e)))
    kv("Hariç tutulan alanlar", "Sistem tarafından otomatik eklenenler (ENA-*, INSDC-* vb.).")
    kv("Alan adı birleştirme", "Küçük harf, boşluk/tire → '_'; ALIASES sözlüğü ve NCBI "
                               "eşanlamlı listesiyle aynı anlamdaki adlar tek ada indirgendi.")
    row += 1

    section("2. Sütunların anlamı")
    for k, v in COLUMN_GUIDE:
        kv(k, v)
    row += 1

    section("3. Renk anahtarı")
    ws.cell(row, 1, "Yüzde sütunları (renk ölçeği)").font = Font(name="Calibri", size=10, bold=True)
    row += 1
    for p in (0, 25, 50, 75, 100):
        lo, hi = (SCALE[0], SCALE[1]) if p <= 50 else (SCALE[1], SCALE[2])
        t = (p - lo[0]) / (hi[0] - lo[0])
        rgb = "".join(f"{round(int(lo[1][i:i+2], 16) * (1 - t) + int(hi[1][i:i+2], 16) * t):02X}"
                      for i in (0, 2, 4))
        c = ws.cell(row, 1, f"%{p}")
        c.fill = PatternFill("solid", fgColor=rgb)
        c.border = border
        c.alignment = Alignment(horizontal="center")
        c.font = Font(name="Calibri", size=10, color="FFFFFF" if p == 100 else "000000")
        if p == 0:
            ws.cell(row, 2, "Beyazdan koyu maviye: alan ne kadar sık doluysa hücre o kadar "
                            "koyu. Tüm yüzde sütunlarında aynı sabit ölçek (0–100) "
                            "kullanılır, bu yüzden sayfalar arası renkler karşılaştırılabilir."
                    ).font = f_body
            ws.cell(row, 2).alignment = wrap_top
        row += 1
    row += 1
    ws.cell(row, 1, "Doluluk sınıfı").font = Font(name="Calibri", size=10, bold=True)
    row += 1
    for _, label, bg, fg in CLASS_BANDS:
        c = ws.cell(row, 1, label)
        c.fill = PatternFill("solid", fgColor=bg)
        c.font = Font(name="Calibri", size=10, color=fg, bold=label.startswith("Yüksek"))
        c.border = border
        row += 1
    row += 1
    legend = [
        ("Gri çubuk", "Sayı sütunlarında (örnek, proje, deney) değerin sütun içindeki büyüklüğü."),
        ("Kalın lacivert", f"kac_kaynakta = {len(sources)}: alan tüm kaynaklarda kullanılıyor."),
        ("İtalik gri", "kac_kaynakta = 1: alan yalnızca tek kaynakta; kaynağa özgü ya da "
                       "standart dışı bir ad olabilir."),
    ]
    for k, v in legend:
        kv(k, v)
    row += 1

    section("4. Yorumlarken dikkat")
    kv("Proje yanlılığı", "Yüksek yüzdeli ama proje_sayisi düşük alanlar tek bir büyük projeyi "
                          "yansıtıyor olabilir; genelleme yapmadan önce kontrol edin.")
    kv("%0 değerleri", "Alan o kaynakta farklı bir adla kayıtlı olabilir. Şüpheli durumlarda "
                       "ALIASES sözlüğünü genişletin.")
    kv("INSDC aynalaması", "NCBI ve ENA verileri büyük ölçüde ortaktır; kaynaklar bağımsız "
                           "örneklemler değildir.")
    kv("Örneklem", f"Her kaynakta en fazla {MAX_RECORDS} kayıt incelendi; tam veri kümesi için "
                   "MAX_RECORDS artırılmalıdır.")

    # ---- 2) Özet
    ws = wb.create_sheet("Ozet")
    ws.append(["Gösterge"] + sources)
    for key, label in SUMMARY_LABELS.items():
        ws.append([label] + [results[s][1][key] for s in sources])
    ws.column_dimensions["A"].width = 38
    for j in range(1, len(sources) + 2):
        L = get_column_letter(j)
        if j > 1:
            ws.column_dimensions[L].width = 14
        for i in range(1, ws.max_row + 1):
            c = ws.cell(i, j)
            c.border = border
            if i == 1:
                c.font, c.fill = f_head, fill_head
                c.alignment = Alignment(horizontal="center")
            else:
                c.font = f_body
                if j > 1:
                    c.alignment = Alignment(horizontal="center")
                    c.number_format = "#,##0.0" if isinstance(c.value, float) else "#,##0"
    for i in range(7, 11):   # doluluk sınıfı satırlarını aynı renklerle işaretle
        _, _, bg, fg = CLASS_BANDS[i - 7]
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=bg)
        ws.cell(i, 1).font = Font(name="Calibri", size=10, color=fg)
    ws.freeze_panes = "B2"

    # ---- 3) Karşılaştırma
    if table:
        cols = ["alan"] + [f"{s}_%" for s in sources] + \
               ["ortalama_%", "kac_kaynakta", "doluluk_sinifi", "aciklama"]
        table_sheet(wb.create_sheet("Karsilastirma"), table, cols,
                    pct_cols={f"{s}_%" for s in sources} | {"ortalama_%"})

    # ---- 4) Kaynak sayfaları
    for s in sources:
        cols = ["alan", "dolu_ornek", "yuzde", "proje_sayisi", "doluluk_sinifi", "aciklama"]
        table_sheet(wb.create_sheet(s), results[s][0], cols, pct_cols={"yuzde"})

    if extras.get("sra_teknik"):
        rows = [dict(r, doluluk_sinifi=doluluk_sinifi(r["yuzde"])) for r in extras["sra_teknik"]]
        table_sheet(wb.create_sheet("SRA_teknik"), rows,
                    ["alan", "dolu_deney", "yuzde", "doluluk_sinifi", "aciklama"], {"yuzde"})

    # ---- 5) Grafik
    if chart:
        try:
            from openpyxl.drawing.image import Image
            ws = wb.create_sheet("Grafik")
            img = Image(chart)
            img.width, img.height = img.width * 0.45, img.height * 0.45
            ws.add_image(img, "B2")
        except Exception as e:
            print(f"(Grafik Excel'e eklenemedi: {e})")

    wb.save("metadata_raporu.xlsx")


def main():
    results, extras = {}, {}
    load_ncbi_dictionary()   # alan açıklamaları + eşanlamlı adlar (parsing'den ÖNCE)

    if RUN_BIOSAMPLE:
        recs = fetch_biosample()
        rows, summ, cl = summarize(recs)
        results["BioSample"] = (rows, summ, cl, recs)

    if RUN_SRA:
        recs, tech, n_exp = fetch_sra()
        rows, summ, cl = summarize(recs)
        results["SRA"] = (rows, summ, cl, recs)
        if n_exp:
            extras["sra_teknik"] = [{"alan": k, "dolu_deney": v,
                                     "yuzde": round(100 * v / n_exp, 1),
                                     "aciklama": describe(k)}
                                    for k, v in tech.most_common()]

    if RUN_ENA:
        recs = fetch_ena()
        rows, summ, cl = summarize(recs)
        results["ENA"] = (rows, summ, cl, recs)

    for s, (rows, summ, cl, _) in results.items():
        extra = None
        if s == "SRA" and "sra_teknik" in extras:
            extra = " Deney düzeyi teknik alanlar: " + ", ".join(
                f"{r['alan'][4:]} %{r['yuzde']:.0f}" for r in extras["sra_teknik"])
        print_source(s, rows, summ, cl, extra)

    results = {s: v for s, v in results.items() if v[1]["ornek"] > 0}
    if len(results) >= 2:
        table = compare(results)
        print_compare(table, list(results))
        print_overlap(results)
    else:
        table = compare(results) if results else []

    if results:
        save_outputs(results, table, extras)


if __name__ == "__main__":
    main()
